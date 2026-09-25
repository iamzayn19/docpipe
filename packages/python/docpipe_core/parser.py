from __future__ import annotations

import mimetypes
from pathlib import Path

from .models import Document, Page, TextBlock


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".xml", ".html"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


class DocpipeError(RuntimeError):
    pass


def parse(path: str | Path, *, ocr: str = "auto") -> Document:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)
    if source.is_dir():
        raise IsADirectoryError(source)

    suffix = source.suffix.lower()
    mime_type = mimetypes.guess_type(source.name)[0] or "application/octet-stream"

    if suffix == ".pdf":
        return _parse_pdf(source, mime_type, ocr=ocr)
    if suffix == ".docx":
        return _parse_docx(source, mime_type)
    if suffix in IMAGE_SUFFIXES:
        return _parse_image(source, mime_type, ocr=ocr)
    if suffix in TEXT_SUFFIXES:
        return _parse_text(source, mime_type)

    return _parse_text(source, mime_type, fallback=True)


def _parse_text(path: Path, mime_type: str, *, fallback: bool = False) -> Document:
    warnings: list[str] = []
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="latin-1")
        warnings.append("Decoded file as latin-1 after utf-8 failed.")
    if fallback:
        warnings.append(f"Unknown extension {path.suffix!r}; parsed as text.")

    page = Page(number=1, text=text, blocks=[TextBlock(text=text, page=1)])
    return Document(
        source=str(path),
        mime_type=mime_type,
        pages=[page],
        metadata={"filename": path.name},
        warnings=warnings,
        backend="text",
    )


def _parse_pdf(path: Path, mime_type: str, *, ocr: str) -> Document:
    try:
        import fitz  # type: ignore[import-not-found]
    except Exception as exc:
        raise DocpipeError(
            "PDF parsing requires PyMuPDF. Install with: pip install 'docpipe-core[pdf]'"
        ) from exc

    pages: list[Page] = []
    warnings: list[str] = []
    with fitz.open(path) as pdf:
        metadata = {k: v for k, v in (pdf.metadata or {}).items() if v}
        for index, raw_page in enumerate(pdf, start=1):
            blocks: list[TextBlock] = []
            for block in raw_page.get_text("blocks"):
                x0, y0, x1, y1, text, *_ = block
                if text.strip():
                    blocks.append(
                        TextBlock(
                            text=text.strip(),
                            page=index,
                            bbox=(float(x0), float(y0), float(x1), float(y1)),
                        )
                    )
            page_text = "\n\n".join(block.text for block in blocks)
            if not page_text and ocr in {"auto", "force"}:
                ocr_text = _ocr_pdf_page(raw_page)
                if ocr_text:
                    page_text = ocr_text
                    blocks.append(TextBlock(text=ocr_text, page=index, kind="ocr"))
                else:
                    warnings.append(f"Page {index} had no extractable text and OCR produced no text.")
            pages.append(Page(number=index, text=page_text, blocks=blocks))

    return Document(
        source=str(path),
        mime_type=mime_type,
        pages=pages,
        metadata={"filename": path.name, **metadata},
        warnings=warnings,
        backend="pymupdf",
    )


def _ocr_pdf_page(page: object) -> str:
    try:
        import fitz  # type: ignore[import-not-found]
        import pytesseract  # type: ignore[import-not-found]
        from PIL import Image  # type: ignore[import-not-found]
    except Exception:
        return ""

    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    mode = "RGBA" if pix.alpha else "RGB"
    image = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
    return pytesseract.image_to_string(image).strip()


def _parse_image(path: Path, mime_type: str, *, ocr: str) -> Document:
    if ocr == "off":
        page = Page(number=1)
        return Document(
            source=str(path),
            mime_type=mime_type,
            pages=[page],
            metadata={"filename": path.name},
            warnings=["OCR disabled; image text was not extracted."],
            backend="image",
        )
    try:
        import pytesseract  # type: ignore[import-not-found]
        from PIL import Image  # type: ignore[import-not-found]
    except Exception as exc:
        raise DocpipeError(
            "Image OCR requires Pillow and pytesseract. Install with: pip install 'docpipe-core[ocr]'"
        ) from exc

    with Image.open(path) as image:
        text = pytesseract.image_to_string(image).strip()
    page = Page(number=1, text=text, blocks=[TextBlock(text=text, page=1, kind="ocr")] if text else [])
    return Document(
        source=str(path),
        mime_type=mime_type,
        pages=[page],
        metadata={"filename": path.name},
        warnings=[] if text else ["OCR produced no text."],
        backend="pytesseract",
    )


def _parse_docx(path: Path, mime_type: str) -> Document:
    try:
        import docx  # type: ignore[import-not-found]
    except Exception as exc:
        raise DocpipeError(
            "DOCX parsing requires python-docx. Install with: pip install 'docpipe-core[docx]'"
        ) from exc

    document = docx.Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    page = Page(number=1, text=text, blocks=[TextBlock(text=t, page=1) for t in paragraphs])
    return Document(
        source=str(path),
        mime_type=mime_type,
        pages=[page],
        metadata={"filename": path.name},
        backend="python-docx",
    )
